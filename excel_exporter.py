# excel_exporter.py
"""Exporta el reporte de una partida a Excel (.xlsx).

openpyxl se importa solo cuando de verdad se exporta, no al cargar el módulo.
Importarlo cuesta ~1,9 s (intenta además cargar numpy), y como los tres
minijuegos importan este módulo, ese tiempo se sumaba al arranque del juego. En
el navegador era peor: pygbag no trae openpyxl, y la pantalla de "Loading,
please wait" no termina hasta que se resuelven los import del juego.

En el navegador ni se intenta: el archivo quedaría en un disco virtual de la
pestaña, inaccesible para el estudiante.
"""
import os
import sys
from datetime import datetime

EN_NAVEGADOR = sys.platform == "emscripten"


def _carpeta_destino():
    descargas = os.path.join(os.path.expanduser('~'), 'Downloads')
    return descargas if os.path.isdir(descargas) else os.getcwd()


def _guardar_con_reintento(wb, carpeta, base):
    """Guarda el archivo. Si está abierto en Excel, Windows lo bloquea y save()
    lanza PermissionError: en ese caso probamos con un nombre alternativo en vez
    de dejar que la excepción tumbe el juego."""
    for intento in range(3):
        sufijo = '' if intento == 0 else f'_{intento + 1}'
        ruta = os.path.join(carpeta, f"{base}{sufijo}.xlsx")
        try:
            wb.save(ruta)
            return ruta, None
        except PermissionError:
            continue
        except Exception as e:
            return None, f"No se pudo guardar: {e}"
    return None, "El archivo está abierto en Excel. Ciérralo e intenta de nuevo."


def export_to_excel(level, score, moves_count, matrix_history, game_events, matrix_stats,
                    resumen_habilidades=None, tema_nombre=None):
    """Exporta la partida a .xlsx. Devuelve (ruta, error); uno de los dos es None.

    matrix_history es una lista de registros {'matriz': [[...]], 'evento': '...'},
    de modo que cada matriz viaja junto a su propio evento y no puede desalinearse.
    """
    if EN_NAVEGADOR:
        return None, "Exportar a Excel solo está disponible en la versión de computador."
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None, "Falta openpyxl. Instálalo con: pip install openpyxl"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen del Juego"

    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

    ws.merge_cells('A1:E1')
    ws['A1'] = "CANDY MATRIX - RESUMEN DEL JUEGO"
    ws['A1'].font = title_font

    info_rows = [
        ("Fecha", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Tema", tema_nombre or "Matrices"),
        ("Nivel alcanzado", level),
        ("Puntuación final", score),
        ("Movimientos totales", moves_count),
    ]
    start_row = 3
    for i, (campo, valor) in enumerate(info_rows, start_row):
        ws.cell(row=i, column=1, value=campo).font = header_font
        ws.cell(row=i, column=1).fill = header_fill
        ws.cell(row=i, column=2, value=valor)

    stats_row = start_row + len(info_rows) + 2
    ws.cell(row=stats_row, column=1, value="ESTADÍSTICAS DE MATRIZ FINAL").font = title_font
    stat_items = [
        ("Suma total", matrix_stats.get('sum', 0)),
        ("Promedio", round(matrix_stats.get('mean', 0), 2)),
        ("Mínimo", matrix_stats.get('min', 0)),
        ("Máximo", matrix_stats.get('max', 0)),
    ]
    for idx, (campo, valor) in enumerate(stat_items, 1):
        ws.cell(row=stats_row + idx, column=1, value=campo)
        ws.cell(row=stats_row + idx, column=2, value=valor)

    if resumen_habilidades:
        fila = stats_row + len(stat_items) + 2
        ws.cell(row=fila, column=1, value="HABILIDADES USADAS").font = title_font
        for idx, (nombre_hab, veces) in enumerate(resumen_habilidades.items(), 1):
            ws.cell(row=fila + idx, column=1, value=nombre_hab)
            ws.cell(row=fila + idx, column=2, value=veces)

    # ---- Hoja 2: historial de matrices, cada una con su evento ----
    ws2 = wb.create_sheet(title="Historial de Matrices")
    ws2.cell(row=1, column=1, value="Jugada").font = header_font
    ws2.cell(row=1, column=2, value="Evento").font = header_font
    fila = 2
    for idx, registro in enumerate(matrix_history, 1):
        if isinstance(registro, dict):
            matriz = registro.get('matriz') or []
            evento = registro.get('evento', '')
        else:
            matriz, evento = registro, ''

        ws2.cell(row=fila, column=1, value=f"Jugada {idx}").font = Font(bold=True)
        ws2.cell(row=fila, column=2, value=evento)
        fila += 1
        for i, mat_row in enumerate(matriz):
            for j, val in enumerate(mat_row):
                ws2.cell(row=fila + i, column=j + 1, value=("#" if val is None else val))
        fila += len(matriz) + 1

    # ---- Hoja 3: registro completo de eventos ----
    ws3 = wb.create_sheet(title="Registro de Eventos")
    ws3.cell(row=1, column=1, value="#").font = header_font
    ws3.cell(row=1, column=2, value="Evento").font = header_font
    for idx, evento in enumerate(game_events, 1):
        ws3.cell(row=idx + 1, column=1, value=idx)
        ws3.cell(row=idx + 1, column=2, value=str(evento))

    for hoja in (ws, ws2, ws3):
        for col in range(1, hoja.max_column + 1):
            letra = get_column_letter(col)
            maxlen = 0
            for celda in hoja[letra]:
                if celda.value is not None:
                    maxlen = max(maxlen, len(str(celda.value)))
            hoja.column_dimensions[letra].width = min(60, maxlen + 4)
    ws3.column_dimensions['B'].alignment = Alignment(wrap_text=True)

    base = f"candy_matrix_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    return _guardar_con_reintento(wb, _carpeta_destino(), base)
